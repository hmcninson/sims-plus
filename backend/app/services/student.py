"""
SIMS Plus - Student Service

Business logic for student management.
"""

import csv
import io
from datetime import datetime, UTC, date
from typing import Optional, Sequence, Any
from uuid import UUID

from sqlalchemy import select, and_, or_, func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.student import (
    Student,
    StudentStatus,
    Gender,
    Guardian,
    StudentGuardian,
    GuardianRelationship,
)
from app.models.academic import Class, ClassSection

# Try to import openpyxl for Excel support
try:
    from openpyxl import load_workbook
    EXCEL_SUPPORT = True
except ImportError:
    EXCEL_SUPPORT = False


class StudentServiceError(Exception):
    """Base exception for student service errors."""

    def __init__(self, message: str, code: str = "student_error"):
        self.message = message
        self.code = code
        super().__init__(message)


class StudentService:
    """Service for managing students and guardians."""

    def __init__(self, db: AsyncSession):
        self.db = db

    # =========================
    # Student Methods
    # =========================

    async def create_student(
        self,
        tenant_id: UUID,
        student_id: str,
        first_name: str,
        last_name: str,
        date_of_birth,
        gender: str,
        middle_name: Optional[str] = None,
        previous_student_id: Optional[str] = None,
        email: Optional[str] = None,
        phone: Optional[str] = None,
        address: Optional[str] = None,
        city: Optional[str] = None,
        region: Optional[str] = None,
        ghana_card_number: Optional[str] = None,
        nhis_number: Optional[str] = None,
        school_id: Optional[UUID] = None,
        class_id: Optional[UUID] = None,
        section_id: Optional[UUID] = None,
        admission_date=None,
        admission_number: Optional[str] = None,
        status: str = "active",
        is_boarder: bool = False,
        blood_group: Optional[str] = None,
        medical_conditions: Optional[str] = None,
        allergies: Optional[str] = None,
        photo_url: Optional[str] = None,
        notes: Optional[str] = None,
        guardians: Optional[list[dict]] = None,
    ) -> Student:
        """Create a new student."""
        # Check for duplicate student_id
        existing = await self.db.execute(
            select(Student).where(
                and_(
                    Student.tenant_id == tenant_id,
                    Student.student_id == student_id,
                    Student.deleted_at.is_(None),
                )
            )
        )
        if existing.scalar_one_or_none():
            raise StudentServiceError(
                f"Student ID '{student_id}' already exists",
                code="duplicate_student_id",
            )

        student = Student(
            tenant_id=tenant_id,
            student_id=student_id,
            previous_student_id=previous_student_id,
            first_name=first_name,
            middle_name=middle_name,
            last_name=last_name,
            date_of_birth=date_of_birth,
            gender=Gender(gender),
            email=email,
            phone=phone,
            address=address,
            city=city,
            region=region,
            ghana_card_number=ghana_card_number,
            nhis_number=nhis_number,
            school_id=school_id,
            class_id=class_id,
            section_id=section_id,
            admission_date=admission_date,
            admission_number=admission_number,
            status=StudentStatus(status),
            is_boarder=is_boarder,
            blood_group=blood_group,
            medical_conditions=medical_conditions,
            allergies=allergies,
            photo_url=photo_url,
            notes=notes,
        )
        self.db.add(student)
        try:
            await self.db.flush()
        except IntegrityError as e:
            await self.db.rollback()
            error_msg = str(e.orig) if e.orig else str(e)
            if "uq_students_student_id_tenant" in error_msg or "student_id" in error_msg.lower():
                raise StudentServiceError(
                    f"Student ID '{student_id}' already exists (concurrent creation detected)",
                    code="duplicate_student_id",
                )
            raise StudentServiceError(
                f"Database error while creating student: {error_msg}",
                code="database_error",
            )

        # Create guardians if provided
        if guardians:
            for guardian_data in guardians:
                # Create the guardian
                guardian = Guardian(
                    tenant_id=tenant_id,
                    first_name=guardian_data["guardian"]["first_name"],
                    last_name=guardian_data["guardian"]["last_name"],
                    phone=guardian_data["guardian"]["phone"],
                    phone_secondary=guardian_data["guardian"].get("phone_secondary"),
                    email=guardian_data["guardian"].get("email"),
                    address=guardian_data["guardian"].get("address"),
                    city=guardian_data["guardian"].get("city"),
                    region=guardian_data["guardian"].get("region"),
                    occupation=guardian_data["guardian"].get("occupation"),
                    workplace=guardian_data["guardian"].get("workplace"),
                    work_phone=guardian_data["guardian"].get("work_phone"),
                    ghana_card_number=guardian_data["guardian"].get("ghana_card_number"),
                    photo_url=guardian_data["guardian"].get("photo_url"),
                    notes=guardian_data["guardian"].get("notes"),
                )
                self.db.add(guardian)
                await self.db.flush()

                # Link guardian to student
                student_guardian = StudentGuardian(
                    tenant_id=tenant_id,
                    student_id=student.id,
                    guardian_id=guardian.id,
                    relation_type=GuardianRelationship(guardian_data["relationship"]),
                    is_primary=guardian_data.get("is_primary", False),
                    is_emergency_contact=guardian_data.get("is_emergency_contact", True),
                    can_pickup=guardian_data.get("can_pickup", True),
                )
                self.db.add(student_guardian)

        await self.db.commit()
        await self.db.refresh(student)
        return student

    async def get_student(
        self, student_id: UUID, include_guardians: bool = False
    ) -> Optional[Student]:
        """Get student by ID."""
        query = select(Student).where(
            and_(
                Student.id == student_id,
                Student.deleted_at.is_(None),
            )
        )
        if include_guardians:
            query = query.options(
                selectinload(Student.guardians).selectinload(StudentGuardian.guardian_rel)
            )
        result = await self.db.execute(query)
        return result.scalar_one_or_none()

    async def get_student_by_student_id(
        self, tenant_id: UUID, student_id: str
    ) -> Optional[Student]:
        """Get student by student ID (the school-assigned ID)."""
        result = await self.db.execute(
            select(Student).where(
                and_(
                    Student.tenant_id == tenant_id,
                    Student.student_id == student_id,
                    Student.deleted_at.is_(None),
                )
            )
        )
        return result.scalar_one_or_none()

    async def list_students(
        self,
        tenant_id: UUID,
        search: Optional[str] = None,
        class_id: Optional[UUID] = None,
        section_id: Optional[UUID] = None,
        school_id: Optional[UUID] = None,
        status: Optional[str] = None,
        gender: Optional[str] = None,
        is_boarder: Optional[bool] = None,
        page: int = 1,
        page_size: int = 20,
    ) -> tuple[Sequence[Student], int]:
        """List students with filters and pagination."""
        conditions = [
            Student.tenant_id == tenant_id,
            Student.deleted_at.is_(None),
        ]

        if search:
            search_term = f"%{search}%"
            conditions.append(
                or_(
                    Student.first_name.ilike(search_term),
                    Student.last_name.ilike(search_term),
                    Student.middle_name.ilike(search_term),
                    Student.student_id.ilike(search_term),
                )
            )

        if class_id:
            conditions.append(Student.class_id == class_id)
        if section_id:
            conditions.append(Student.section_id == section_id)
        if school_id:
            conditions.append(Student.school_id == school_id)
        if status:
            conditions.append(Student.status == StudentStatus(status))
        if gender:
            conditions.append(Student.gender == Gender(gender))
        if is_boarder is not None:
            conditions.append(Student.is_boarder == is_boarder)

        # Get total count
        count_query = select(func.count(Student.id)).where(and_(*conditions))
        total = (await self.db.execute(count_query)).scalar_one()

        # Get paginated results
        query = (
            select(Student)
            .where(and_(*conditions))
            .options(
                selectinload(Student.class_),
                selectinload(Student.section),
            )
            .order_by(Student.last_name, Student.first_name)
            .offset((page - 1) * page_size)
            .limit(page_size)
        )
        result = await self.db.execute(query)
        students = result.scalars().all()

        return students, total

    async def update_student(
        self, student_id: UUID, **kwargs
    ) -> Optional[Student]:
        """Update a student."""
        student = await self.get_student(student_id)
        if not student:
            return None

        # Handle enum conversions
        if "status" in kwargs and kwargs["status"]:
            kwargs["status"] = StudentStatus(kwargs["status"])
        if "gender" in kwargs and kwargs["gender"]:
            kwargs["gender"] = Gender(kwargs["gender"])

        for key, value in kwargs.items():
            if value is not None and hasattr(student, key):
                setattr(student, key, value)

        student.updated_at = datetime.now(UTC)
        await self.db.commit()
        await self.db.refresh(student)
        return student

    async def delete_student(self, student_id: UUID) -> bool:
        """Soft delete a student."""
        student = await self.get_student(student_id)
        if not student:
            return False

        student.deleted_at = datetime.now(UTC)
        await self.db.commit()
        return True

    async def generate_student_id(
        self, tenant_id: UUID, prefix: str = "STU", max_retries: int = 10
    ) -> str:
        """
        Generate a unique student ID for the tenant.

        Format: {PREFIX}-{YEAR}-{SEQUENCE}
        Example: STU-2026-001, STU-2026-002

        This method finds the highest existing sequence number for the given
        prefix and year pattern, then increments it. It includes retry logic
        for race condition protection.

        Args:
            tenant_id: The tenant UUID
            prefix: The student ID prefix (default: STU)
            max_retries: Maximum number of retry attempts for race conditions

        Returns:
            A unique student ID string

        Raises:
            StudentServiceError: If unable to generate unique ID after max retries
        """
        import re

        current_year = datetime.now().year
        pattern = f"{prefix}-{current_year}-%"

        # Find the highest existing sequence number for this prefix and year
        # This query finds all student IDs matching the pattern and extracts the max sequence
        result = await self.db.execute(
            select(Student.student_id).where(
                and_(
                    Student.tenant_id == tenant_id,
                    Student.student_id.like(pattern),
                )
            )
        )
        existing_ids = result.scalars().all()

        # Extract sequence numbers from existing IDs
        max_seq = 0
        seq_pattern = re.compile(rf"^{re.escape(prefix)}-{current_year}-(\d+)$")
        for existing_id in existing_ids:
            match = seq_pattern.match(existing_id)
            if match:
                seq = int(match.group(1))
                if seq > max_seq:
                    max_seq = seq

        # Start from the next sequence number
        next_seq = max_seq + 1

        # Retry loop for race condition protection
        for attempt in range(max_retries):
            # Format: PREFIX-YEAR-SEQUENCE (3 digits minimum, auto-expand if needed)
            if next_seq < 1000:
                student_id = f"{prefix}-{current_year}-{next_seq:03d}"
            else:
                # Support larger schools with more than 999 students
                student_id = f"{prefix}-{current_year}-{next_seq}"

            # Check if this ID already exists (including soft-deleted for safety)
            existing = await self.db.execute(
                select(Student.id).where(
                    and_(
                        Student.tenant_id == tenant_id,
                        Student.student_id == student_id,
                    )
                )
            )
            if not existing.scalar_one_or_none():
                return student_id

            # ID exists, try next sequence number
            next_seq += 1

        # If we've exhausted all retries, raise an error
        raise StudentServiceError(
            f"Unable to generate unique student ID after {max_retries} attempts. "
            f"Please try again or assign a student ID manually.",
            code="id_generation_failed",
        )

    async def get_student_stats(self, tenant_id: UUID) -> dict:
        """Get student statistics for a tenant."""
        base_conditions = [
            Student.tenant_id == tenant_id,
            Student.deleted_at.is_(None),
        ]

        # Total count
        total = (
            await self.db.execute(
                select(func.count(Student.id)).where(and_(*base_conditions))
            )
        ).scalar_one()

        # By status
        status_counts = {}
        for status in StudentStatus:
            count = (
                await self.db.execute(
                    select(func.count(Student.id)).where(
                        and_(*base_conditions, Student.status == status)
                    )
                )
            ).scalar_one()
            status_counts[status.value] = count

        # By gender
        male_count = (
            await self.db.execute(
                select(func.count(Student.id)).where(
                    and_(*base_conditions, Student.gender == Gender.MALE)
                )
            )
        ).scalar_one()

        female_count = (
            await self.db.execute(
                select(func.count(Student.id)).where(
                    and_(*base_conditions, Student.gender == Gender.FEMALE)
                )
            )
        ).scalar_one()

        # Boarders vs day students
        boarder_count = (
            await self.db.execute(
                select(func.count(Student.id)).where(
                    and_(*base_conditions, Student.is_boarder == True)
                )
            )
        ).scalar_one()

        return {
            "total": total,
            "active": status_counts.get("active", 0),
            "inactive": status_counts.get("inactive", 0),
            "graduated": status_counts.get("graduated", 0),
            "transferred": status_counts.get("transferred", 0),
            "withdrawn": status_counts.get("withdrawn", 0),
            "suspended": status_counts.get("suspended", 0),
            "male": male_count,
            "female": female_count,
            "boarders": boarder_count,
            "day_students": total - boarder_count,
        }

    async def bulk_create_students(
        self,
        tenant_id: UUID,
        students_data: list[dict],
    ) -> dict:
        """Bulk create students."""
        created = 0
        failed = 0
        errors = []

        for i, student_data in enumerate(students_data):
            try:
                await self.create_student(
                    tenant_id=tenant_id,
                    **student_data,
                )
                created += 1
            except StudentServiceError as e:
                failed += 1
                errors.append({
                    "index": i,
                    "student_id": student_data.get("student_id"),
                    "error": e.message,
                })
            except Exception as e:
                failed += 1
                errors.append({
                    "index": i,
                    "student_id": student_data.get("student_id"),
                    "error": str(e),
                })

        return {
            "created": created,
            "failed": failed,
            "errors": errors,
        }

    # =========================
    # File Import Methods
    # =========================

    @staticmethod
    def parse_csv_file(file_content: bytes) -> tuple[list[str], list[dict]]:
        """
        Parse a CSV file and return headers and rows.

        Args:
            file_content: Raw bytes of the CSV file

        Returns:
            Tuple of (headers, rows as list of dicts)
        """
        # Try different encodings
        content = None
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                content = file_content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if content is None:
            raise StudentServiceError("Unable to decode file. Please use UTF-8 encoding.")

        reader = csv.DictReader(io.StringIO(content))
        headers = reader.fieldnames or []
        rows = list(reader)

        return headers, rows

    @staticmethod
    def parse_excel_file(file_content: bytes) -> tuple[list[str], list[dict]]:
        """
        Parse an Excel file and return headers and rows.

        Args:
            file_content: Raw bytes of the Excel file

        Returns:
            Tuple of (headers, rows as list of dicts)
        """
        if not EXCEL_SUPPORT:
            raise StudentServiceError("Excel support not available. Install openpyxl.")

        workbook = load_workbook(filename=io.BytesIO(file_content), read_only=True)
        sheet = workbook.active

        rows = []
        headers = []

        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0:
                # First row is headers
                headers = [str(cell).strip() if cell else f"Column_{j}" for j, cell in enumerate(row)]
            else:
                # Convert row to dict
                row_dict = {}
                for j, cell in enumerate(row):
                    if j < len(headers):
                        row_dict[headers[j]] = cell
                # Skip empty rows
                if any(v is not None and v != "" for v in row_dict.values()):
                    rows.append(row_dict)

        workbook.close()
        return headers, rows

    @staticmethod
    def _normalize_column_name(name: str) -> str:
        """Normalize column name for matching."""
        return name.lower().strip().replace(" ", "_").replace("-", "_")

    @staticmethod
    def _auto_map_columns(headers: list[str]) -> dict[str, str]:
        """
        Auto-map CSV/Excel headers to student fields.

        Returns:
            Dict mapping student field names to CSV column names
        """
        # Map file columns to our fields
        # Note: student_id from file is mapped to previous_student_id
        # The system always auto-generates the actual student_id
        field_aliases = {
            "previous_student_id": ["student_id", "studentid", "student_no", "student_number", "id", "admission_no", "previous_student_id", "old_id", "external_id"],
            "first_name": ["first_name", "firstname", "first", "given_name"],
            "middle_name": ["middle_name", "middlename", "middle", "other_names"],
            "last_name": ["last_name", "lastname", "last", "surname", "family_name"],
            "date_of_birth": ["date_of_birth", "dob", "dateofbirth", "birth_date", "birthdate"],
            "gender": ["gender", "sex"],
            "email": ["email", "email_address", "e_mail"],
            "phone": ["phone", "phone_number", "telephone", "mobile", "contact"],
            "address": ["address", "home_address", "residential_address"],
            "city": ["city", "town"],
            "region": ["region", "state", "province"],
            "class_name": ["class", "class_name", "grade", "level", "form"],
            "is_boarder": ["is_boarder", "boarder", "boarding", "resident"],
        }

        mapping = {}
        normalized_headers = {StudentService._normalize_column_name(h): h for h in headers}

        for field, aliases in field_aliases.items():
            for alias in aliases:
                if alias in normalized_headers:
                    mapping[field] = normalized_headers[alias]
                    break

        return mapping

    @staticmethod
    def _parse_date(value: Any) -> Optional[date]:
        """Parse various date formats."""
        if value is None:
            return None
        if isinstance(value, date):
            return value
        if isinstance(value, datetime):
            return value.date()

        value_str = str(value).strip()
        if not value_str:
            return None

        # Try various date formats
        formats = [
            "%Y-%m-%d",      # 2026-01-15
            "%d/%m/%Y",      # 15/01/2026
            "%m/%d/%Y",      # 01/15/2026
            "%d-%m-%Y",      # 15-01-2026
            "%Y/%m/%d",      # 2026/01/15
            "%d %b %Y",      # 15 Jan 2026
            "%d %B %Y",      # 15 January 2026
        ]

        for fmt in formats:
            try:
                return datetime.strptime(value_str, fmt).date()
            except ValueError:
                continue

        raise ValueError(f"Unable to parse date: {value_str}")

    @staticmethod
    def _parse_gender(value: Any) -> str:
        """Parse gender value to standard format."""
        if value is None:
            return "male"  # Default

        value_str = str(value).lower().strip()
        if value_str in ["m", "male", "boy", "man"]:
            return "male"
        elif value_str in ["f", "female", "girl", "woman"]:
            return "female"
        else:
            return "male"  # Default

    @staticmethod
    def _parse_boolean(value: Any) -> bool:
        """Parse boolean value."""
        if value is None:
            return False
        if isinstance(value, bool):
            return value

        value_str = str(value).lower().strip()
        return value_str in ["true", "yes", "1", "y", "t", "boarder"]

    async def import_students_from_file(
        self,
        tenant_id: UUID,
        file_content: bytes,
        file_type: str,
        column_mapping: Optional[dict[str, str]] = None,
        preview_only: bool = False,
        auto_generate_ids: bool = True,  # Kept for backward compatibility, but always generates
    ) -> dict:
        """
        Import students from a CSV or Excel file.

        Args:
            tenant_id: Tenant UUID
            file_content: Raw file bytes
            file_type: 'csv' or 'xlsx'
            column_mapping: Optional manual column mapping
            preview_only: If True, only return preview without creating
            auto_generate_ids: Deprecated - student_id is always auto-generated.
                               File's student_id column is stored as previous_student_id.

        Returns:
            Dict with import results
        """
        # Parse file based on type
        if file_type == "csv":
            headers, rows = self.parse_csv_file(file_content)
        elif file_type in ["xlsx", "xls"]:
            headers, rows = self.parse_excel_file(file_content)
        else:
            raise StudentServiceError(f"Unsupported file type: {file_type}")

        if not rows:
            return {
                "total_rows": 0,
                "created": 0,
                "failed": 0,
                "errors": [{"row": 0, "error": "File is empty or has no data rows"}],
                "preview": [],
            }

        # Auto-map columns if not provided
        if not column_mapping:
            column_mapping = self._auto_map_columns(headers)

        # Preview mode - return first 5 rows with parsed values
        if preview_only:
            preview = []
            for i, row in enumerate(rows[:5]):
                parsed = self._parse_row(row, column_mapping, i + 2)
                preview.append({
                    "row_number": i + 2,
                    "raw": row,
                    "parsed": parsed.get("data", {}),
                    "errors": parsed.get("errors", []),
                })
            return {
                "total_rows": len(rows),
                "created": 0,
                "failed": 0,
                "errors": [],
                "preview": preview,
                "headers": headers,
                "auto_mapping": column_mapping,
            }

        # Import mode - create students
        # Note: student_id is ALWAYS auto-generated by the system
        # The file's student_id column is mapped to previous_student_id
        created = 0
        failed = 0
        errors = []
        max_id_retries = 5  # Max retries for ID conflicts

        # Get the school's student_id_prefix for consistent ID generation
        from app.models.school import School
        school_result = await self.db.execute(
            select(School.student_id_prefix).where(School.tenant_id == tenant_id).limit(1)
        )
        school_prefix = school_result.scalar_one_or_none() or "STU"

        for i, row in enumerate(rows):
            row_number = i + 2  # Account for header row
            parsed = self._parse_row(row, column_mapping, row_number)

            if parsed.get("errors"):
                failed += 1
                errors.extend(parsed["errors"])
                continue

            student_data = parsed["data"]

            # Create the student with retry logic for auto-generated IDs
            student_created = False
            for attempt in range(max_id_retries):
                try:
                    # Always auto-generate student ID using school's prefix
                    student_data["student_id"] = await self.generate_student_id(
                        tenant_id, prefix=school_prefix
                    )

                    await self.create_student(
                        tenant_id=tenant_id,
                        **student_data,
                    )
                    created += 1
                    student_created = True
                    break
                except StudentServiceError as e:
                    # Retry on duplicate ID (race condition on auto-generation)
                    if e.code == "duplicate_student_id" and attempt < max_id_retries - 1:
                        import asyncio
                        await asyncio.sleep(0.01 * (attempt + 1))
                        continue
                    failed += 1
                    errors.append({
                        "row": row_number,
                        "previous_student_id": student_data.get("previous_student_id"),
                        "error": e.message,
                    })
                    break
                except Exception as e:
                    failed += 1
                    errors.append({
                        "row": row_number,
                        "previous_student_id": student_data.get("previous_student_id"),
                        "error": str(e),
                    })
                    break

            if not student_created and not any(
                err.get("row") == row_number for err in errors
            ):
                # All retries exhausted without recording an error
                failed += 1
                errors.append({
                    "row": row_number,
                    "error": f"Failed to generate unique student ID after {max_id_retries} attempts",
                })

        return {
            "total_rows": len(rows),
            "created": created,
            "failed": failed,
            "errors": errors,
            "preview": [],
        }

    def _parse_row(
        self,
        row: dict,
        column_mapping: dict[str, str],
        row_number: int,
    ) -> dict:
        """
        Parse a single row using the column mapping.

        Returns:
            Dict with 'data' and 'errors' keys
        """
        data = {}
        errors = []

        # Required fields
        required_fields = ["first_name", "last_name", "date_of_birth", "gender"]

        for field, column in column_mapping.items():
            if column not in row:
                continue

            value = row[column]

            try:
                if field == "date_of_birth":
                    parsed_value = self._parse_date(value)
                    if parsed_value:
                        data[field] = parsed_value
                elif field == "gender":
                    data[field] = self._parse_gender(value)
                elif field == "is_boarder":
                    data[field] = self._parse_boolean(value)
                elif field == "class_name":
                    # Store for later lookup, don't add to data directly
                    pass
                elif value is not None and str(value).strip():
                    data[field] = str(value).strip()
            except Exception as e:
                errors.append({
                    "row": row_number,
                    "field": field,
                    "value": str(value),
                    "error": str(e),
                })

        # Check required fields
        for field in required_fields:
            if field not in data:
                mapped_column = column_mapping.get(field, field)
                errors.append({
                    "row": row_number,
                    "field": field,
                    "error": f"Required field '{mapped_column}' is missing or empty",
                })

        return {"data": data, "errors": errors}

    # =========================
    # Guardian Methods
    # =========================

    async def create_guardian(
        self,
        tenant_id: UUID,
        first_name: str,
        last_name: str,
        phone: str,
        phone_secondary: Optional[str] = None,
        email: Optional[str] = None,
        address: Optional[str] = None,
        city: Optional[str] = None,
        region: Optional[str] = None,
        occupation: Optional[str] = None,
        workplace: Optional[str] = None,
        work_phone: Optional[str] = None,
        ghana_card_number: Optional[str] = None,
        photo_url: Optional[str] = None,
        notes: Optional[str] = None,
    ) -> Guardian:
        """Create a new guardian."""
        # Check for duplicate email if provided
        if email:
            existing = await self.db.execute(
                select(Guardian).where(
                    and_(
                        Guardian.tenant_id == tenant_id,
                        Guardian.email == email,
                        Guardian.deleted_at.is_(None),
                    )
                )
            )
            if existing.scalar_one_or_none():
                raise StudentServiceError(
                    f"A guardian with email '{email}' already exists",
                    code="duplicate_guardian_email",
                )

        guardian = Guardian(
            tenant_id=tenant_id,
            first_name=first_name,
            last_name=last_name,
            phone=phone,
            phone_secondary=phone_secondary,
            email=email,
            address=address,
            city=city,
            region=region,
            occupation=occupation,
            workplace=workplace,
            work_phone=work_phone,
            ghana_card_number=ghana_card_number,
            photo_url=photo_url,
            notes=notes,
        )
        self.db.add(guardian)
        await self.db.commit()
        await self.db.refresh(guardian)
        return guardian

    async def get_guardian(self, guardian_id: UUID) -> Optional[Guardian]:
        """Get guardian by ID."""
        result = await self.db.execute(
            select(Guardian).where(
                and_(
                    Guardian.id == guardian_id,
                    Guardian.deleted_at.is_(None),
                )
            )
        )
        return result.scalar_one_or_none()

    async def list_guardians(
        self,
        tenant_id: UUID,
        search: Optional[str] = None,
        page: int = 1,
        page_size: int = 20,
    ) -> tuple[list[dict], int]:
        """List guardians with pagination and student counts."""
        conditions = [
            Guardian.tenant_id == tenant_id,
            Guardian.deleted_at.is_(None),
        ]

        if search:
            search_term = f"%{search}%"
            conditions.append(
                or_(
                    Guardian.first_name.ilike(search_term),
                    Guardian.last_name.ilike(search_term),
                    Guardian.phone.ilike(search_term),
                    Guardian.email.ilike(search_term),
                )
            )

        # Get total count
        count_query = select(func.count(Guardian.id)).where(and_(*conditions))
        total = (await self.db.execute(count_query)).scalar_one()

        # Subquery for student count
        student_count_subq = (
            select(func.count(StudentGuardian.id))
            .where(StudentGuardian.guardian_id == Guardian.id)
            .correlate(Guardian)
            .scalar_subquery()
        )

        # Get paginated results with student count
        query = (
            select(Guardian, student_count_subq.label("student_count"))
            .where(and_(*conditions))
            .order_by(Guardian.last_name, Guardian.first_name)
            .offset((page - 1) * page_size)
            .limit(page_size)
        )
        result = await self.db.execute(query)
        rows = result.all()

        # Convert to list of dicts with student_count
        guardians_with_counts = []
        for row in rows:
            guardian = row[0]
            student_count = row[1] or 0
            guardian_dict = {
                "id": guardian.id,
                "first_name": guardian.first_name,
                "last_name": guardian.last_name,
                "phone": guardian.phone,
                "phone_secondary": guardian.phone_secondary,
                "email": guardian.email,
                "address": guardian.address,
                "city": guardian.city,
                "region": guardian.region,
                "occupation": guardian.occupation,
                "workplace": guardian.workplace,
                "work_phone": guardian.work_phone,
                "ghana_card_number": guardian.ghana_card_number,
                "photo_url": guardian.photo_url,
                "notes": guardian.notes,
                "created_at": guardian.created_at,
                "updated_at": guardian.updated_at,
                "student_count": student_count,
            }
            guardians_with_counts.append(guardian_dict)

        return guardians_with_counts, total

    async def update_guardian(
        self, guardian_id: UUID, **kwargs
    ) -> Optional[Guardian]:
        """Update a guardian."""
        guardian = await self.get_guardian(guardian_id)
        if not guardian:
            return None

        for key, value in kwargs.items():
            if value is not None and hasattr(guardian, key):
                setattr(guardian, key, value)

        guardian.updated_at = datetime.now(UTC)
        await self.db.commit()
        await self.db.refresh(guardian)
        return guardian

    async def get_guardian_linked_students_count(self, guardian_id: UUID) -> int:
        """Get the count of students linked to a guardian."""
        result = await self.db.execute(
            select(func.count(StudentGuardian.id)).where(
                StudentGuardian.guardian_id == guardian_id
            )
        )
        return result.scalar() or 0

    async def delete_guardian(self, guardian_id: UUID) -> bool:
        """
        Soft delete a guardian.

        Raises:
            StudentServiceError: If guardian has linked students
        """
        guardian = await self.get_guardian(guardian_id)
        if not guardian:
            return False

        # Check if guardian has linked students
        linked_count = await self.get_guardian_linked_students_count(guardian_id)
        if linked_count > 0:
            raise StudentServiceError(
                f"Cannot delete guardian with {linked_count} linked student(s). "
                "Please unlink all students first.",
                code="guardian_has_linked_students"
            )

        guardian.deleted_at = datetime.now(UTC)
        await self.db.commit()
        return True

    # =========================
    # Student-Guardian Link Methods
    # =========================

    async def link_guardian_to_student(
        self,
        tenant_id: UUID,
        student_id: UUID,
        guardian_id: UUID,
        relationship: str,
        is_primary: bool = False,
        is_emergency_contact: bool = True,
        can_pickup: bool = True,
    ) -> StudentGuardian:
        """Link an existing guardian to a student."""
        # Verify student and guardian exist
        student = await self.get_student(student_id)
        guardian = await self.get_guardian(guardian_id)

        if not student:
            raise StudentServiceError("Student not found", code="student_not_found")
        if not guardian:
            raise StudentServiceError("Guardian not found", code="guardian_not_found")

        # Check for existing link
        existing = await self.db.execute(
            select(StudentGuardian).where(
                and_(
                    StudentGuardian.tenant_id == tenant_id,
                    StudentGuardian.student_id == student_id,
                    StudentGuardian.guardian_id == guardian_id,
                )
            )
        )
        if existing.scalar_one_or_none():
            raise StudentServiceError(
                "Guardian is already linked to this student",
                code="duplicate_link",
            )

        # If setting as primary, unset other primary guardians
        if is_primary:
            await self._unset_primary_guardian(tenant_id, student_id)

        link = StudentGuardian(
            tenant_id=tenant_id,
            student_id=student_id,
            guardian_id=guardian_id,
            relation_type=GuardianRelationship(relationship),
            is_primary=is_primary,
            is_emergency_contact=is_emergency_contact,
            can_pickup=can_pickup,
        )
        self.db.add(link)
        await self.db.commit()
        await self.db.refresh(link)
        return link

    async def update_student_guardian_link(
        self,
        link_id: UUID,
        tenant_id: UUID,
        student_id: UUID,
        **kwargs,
    ) -> Optional[StudentGuardian]:
        """Update a student-guardian link."""
        result = await self.db.execute(
            select(StudentGuardian).where(StudentGuardian.id == link_id)
        )
        link = result.scalar_one_or_none()
        if not link:
            return None

        # Handle relationship enum conversion (API uses 'relationship', model uses 'relation_type')
        if "relationship" in kwargs and kwargs["relationship"]:
            kwargs["relation_type"] = GuardianRelationship(kwargs["relationship"])
            del kwargs["relationship"]

        # If setting as primary, unset other primary guardians
        if kwargs.get("is_primary"):
            await self._unset_primary_guardian(tenant_id, student_id)

        for key, value in kwargs.items():
            if value is not None and hasattr(link, key):
                setattr(link, key, value)

        link.updated_at = datetime.now(UTC)
        await self.db.commit()
        await self.db.refresh(link)
        return link

    async def unlink_guardian_from_student(
        self, student_id: UUID, guardian_id: UUID
    ) -> bool:
        """Remove a guardian link from a student."""
        result = await self.db.execute(
            select(StudentGuardian).where(
                and_(
                    StudentGuardian.student_id == student_id,
                    StudentGuardian.guardian_id == guardian_id,
                )
            )
        )
        link = result.scalar_one_or_none()
        if not link:
            return False

        await self.db.delete(link)
        await self.db.commit()
        return True

    async def get_student_guardians(
        self, tenant_id: UUID, student_id: UUID
    ) -> Sequence[StudentGuardian]:
        """Get all guardians linked to a student."""
        result = await self.db.execute(
            select(StudentGuardian)
            .where(
                and_(
                    StudentGuardian.tenant_id == tenant_id,
                    StudentGuardian.student_id == student_id,
                )
            )
            .options(selectinload(StudentGuardian.guardian_rel))
        )
        return result.scalars().all()

    async def get_guardian_students(
        self, tenant_id: UUID, guardian_id: UUID
    ) -> Sequence[StudentGuardian]:
        """Get all students linked to a guardian."""
        result = await self.db.execute(
            select(StudentGuardian)
            .where(
                and_(
                    StudentGuardian.tenant_id == tenant_id,
                    StudentGuardian.guardian_id == guardian_id,
                )
            )
            .options(selectinload(StudentGuardian.student_rel))
        )
        return result.scalars().all()

    async def _unset_primary_guardian(
        self, tenant_id: UUID, student_id: UUID
    ) -> None:
        """Unset primary flag for all guardians of a student."""
        result = await self.db.execute(
            select(StudentGuardian).where(
                and_(
                    StudentGuardian.tenant_id == tenant_id,
                    StudentGuardian.student_id == student_id,
                    StudentGuardian.is_primary == True,
                )
            )
        )
        for link in result.scalars().all():
            link.is_primary = False
